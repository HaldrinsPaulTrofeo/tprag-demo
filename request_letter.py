"""
request_letter.py — generate the anti-rabies vaccine request letter to the
Office of the Provincial Veterinarian, populated from the session forecast.

WHY THIS MODULE IS THE POINT OF THE FORECAST
--------------------------------------------
The Municipal Agriculture Office does not control how much vaccine it
receives. Allocation is decided by the Provincial Veterinarian, annually,
against a request letter, and gated by liquidation of the previous release
(Session 2 interview; Mr. Caballes, Sept 2026).

That request letter is the one document in the chain the office fully
controls, and it needs exactly two things: a vial count per barangay and a
schedule of dates. Both are what session_forecast.py produces.

So the forecast is not predicting an outcome the office is powerless over. It
is generating the justification for the document that determines the ceiling.

The observed 2026 letter requested 251 vials across seven barangays. Measured
against what those sessions actually consumed, the estimates were off in both
directions — Magdapio requested 60 vials and used 14 (23%); Sabang requested
25 and used 33 (132%). That spread is the gap this module closes.

FORMAT
------
Mirrors the letter the office already sends (addressed to the Provincial
Veterinarian, dated schedule lines reading "70vials Brgy.Binan", signed by the
OIC-Municipal Agriculturist). Keeping the familiar format matters for
adoption: the office can send the generated letter as-is.

USAGE
    python request_letter.py --barangay Binan:2026-03-24 --barangay Sabang:2026-03-26
    python request_letter.py --from-forecast Binan Sabang Lambac --start 2026-03-18
"""

from __future__ import annotations

import argparse
from datetime import date, datetime, timedelta
from pathlib import Path

from session_forecast import DOSES_PER_VIAL, forecast_session, load_sessions

PROVINCIAL_VET = "MICHAEL L. CORTEZ, DVM"
PROVINCIAL_VET_TITLE = "Provincial Veterinarian"
PROVINCE = "Province of Laguna"
SIGNATORY = "REIMEL P. CABALLES RAgr."
SIGNATORY_TITLE = "OIC-Municipal Agriculturist"

BODY = ("Good day! As part of our program in the implementation of Anti-Rabies "
        "vaccination of the different barangay in our municipality the "
        "following date as schedule;")


def plan(barangays, service_level=0.75, start=None, spacing_days=1):
    """Build the vial plan. Returns (rows, total_vials).

    Each row: (barangay, session_date, vials, median_doses, basis, n_sessions).
    """
    sessions = load_sessions()
    rows = []
    cursor = start
    for entry in barangays:
        if ":" in entry:
            name, d = entry.split(":", 1)
            when = datetime.strptime(d.strip(), "%Y-%m-%d").date()
        else:
            name, when = entry, cursor
            if cursor is not None:
                cursor = cursor + timedelta(days=spacing_days)
        f = forecast_session(name.strip(), service_level, sessions)
        rows.append((f.barangay, when, f.recommended_vials,
                     f.median_doses, f.basis, f.n_sessions_barangay))
    return rows, sum(r[2] for r in rows)


def plan_from_requests(requests):
    """Use original requested quantities, never grants or supply-capped recommendations."""
    rows = []
    for request in requests:
        if request["category"] != "ANTI_RABIES" or request["unit"].lower() != "vials":
            raise ValueError("Vaccine letters require anti-rabies requests in vials")
        when = request.get("needed_by")
        if isinstance(when, str):
            when = date.fromisoformat(when)
        rows.append((request["barangay"], when,
                     int(request["qty_requested"]), None, "barangay request", None))
    return rows, sum(row[2] for row in rows)


def write_docx(rows, total_vials, path, service_level):
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(1)
        s.left_margin = s.right_margin = Inches(1)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def p(text="", bold=False, align=None, after=6):
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(after)
        if align:
            par.alignment = align
        r = par.add_run(text)
        r.bold = bold
        return par

    p("Republic of the Philippines", align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    p(PROVINCE, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    p("Municipality of Pagsanjan", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, after=0)
    p("Office of the Municipal Agriculturist", align=WD_ALIGN_PARAGRAPH.CENTER, after=18)

    p(PROVINCIAL_VET, bold=True, after=0)
    p(PROVINCIAL_VET_TITLE, after=0)
    p(PROVINCE, after=18)

    p("Sir,", after=6)
    par = doc.add_paragraph()
    par.paragraph_format.first_line_indent = Inches(0.4)
    par.paragraph_format.space_after = Pt(18)
    par.add_run(BODY)

    for brgy, when, vials, *_ in rows:
        line = doc.add_paragraph()
        line.paragraph_format.space_after = Pt(14)
        d = when.strftime("%B %d,%Y") if when else "____________"
        line.add_run(f"{d} ")
        line.add_run("-" * 38 + " ")
        line.add_run(f"{vials}vials Brgy.{brgy}")

    p(after=24)
    p(SIGNATORY, bold=True, after=0)
    p(SIGNATORY_TITLE, after=0)

    doc.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--barangay", action="append", default=[],
                    help="Barangay or Barangay:YYYY-MM-DD; repeatable")
    ap.add_argument("--from-forecast", nargs="*", default=[],
                    help="barangay names, dates assigned from --start")
    ap.add_argument("--start", help="first session date, YYYY-MM-DD")
    ap.add_argument("--spacing", type=int, default=1, help="days between sessions")
    ap.add_argument("--service-level", type=float, default=0.75)
    ap.add_argument("--out", default="request_letter.docx")
    args = ap.parse_args()

    names = args.barangay + list(args.from_forecast)
    if not names:
        ap.error("give at least one --barangay or --from-forecast")

    start = datetime.strptime(args.start, "%Y-%m-%d").date() if args.start else None
    rows, total = plan(names, args.service_level, start, args.spacing)

    print(f"\n  VIAL PLAN  (service level {args.service_level:.0%})")
    print(f"  {'barangay':<16}{'date':<14}{'vials':>6}{'doses':>7}"
          f"{'typical':>9}  basis")
    for brgy, when, vials, med, basis, nb in rows:
        d = when.isoformat() if when else "—"
        note = f"{basis} (n={nb})" if basis == "barangay" else f"{basis}"
        print(f"  {brgy:<16}{d:<14}{vials:>6}{vials * DOSES_PER_VIAL:>7}"
              f"{med:>9.0f}  {note}")
    print(f"  {'TOTAL':<16}{'':<14}{total:>6}{total * DOSES_PER_VIAL:>7}")

    out = write_docx(rows, total, args.out, args.service_level)
    print(f"\n  wrote {out}")
    print("  Review the dates and vial counts before sending. The forecast is a "
          "planning\n  range, not a prediction — session size depends on crew "
          "capacity and weather.")


if __name__ == "__main__":
    main()


# ---------------------------------------------------------------------------
# Printable letters to the Provincial Veterinarian, composed on the staff
# request page. Download/print only: there is no online exchange with the
# province until the partner office authorises one.
# ---------------------------------------------------------------------------

STATIC = Path(__file__).resolve().parent / "static"
GENERAL_SUBJECT = "Request for Anti-Rabies Vaccination"


def schedule_date(value):
    """'March 24,2026', the spelling used in the office's own letters."""
    return value.strftime("%B %d,%Y").replace(" 0", " ") if value else "____________"


def build_provincial_docx(letter, signature_path=None):
    """Return .docx bytes for a 'schedule' or 'general' letter (see provincial_letter_api)."""
    from io import BytesIO
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.shared import Inches, Pt

    doc = Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Inches(0.6), Inches(0.8)
        s.left_margin = s.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(12)
    normal.paragraph_format.space_after = Pt(0)

    head = doc.add_table(rows=1, cols=3)
    head.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (Inches(1.1), Inches(4.3), Inches(1.1))
    for cell, w in zip(head.rows[0].cells, widths):
        cell.width = w
    for idx, name in ((0, "seal-pagsanjan.png"), (2, "seal-mao.png")):
        if (STATIC / name).exists():
            par = head.rows[0].cells[idx].paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.add_run().add_picture(str(STATIC / name), width=Inches(0.95))
    mid = head.rows[0].cells[1]
    mid.paragraphs[0].text = ""
    for i, line in enumerate(("Republic of the Philippines", "Province of Laguna",
                              "Municipality of Pagsanjan", "Office of the Municipal Agriculturist")):
        par = mid.paragraphs[0] if i == 0 else mid.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = par.add_run(line)
        run.font.name, run.font.size = "Cambria", Pt(12)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(18)
    rule.add_run("_" * 86).font.size = Pt(8)

    def par(text="", bold=False, after=0, indent=False, align=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.4)
        if align:
            p.alignment = align
        if text:
            p.add_run(text).bold = bold
        return p

    par(letter["letter_date"].strftime("%B %d, %Y").replace(" 0", " "), after=18)
    par(letter["addressee_name"], bold=True)
    par(letter["addressee_title"])
    if letter.get("addressee_office"):
        par(letter["addressee_office"])
    par(after=14)
    if letter["template"] == "general":
        par("Subject: " + (letter.get("subject") or GENERAL_SUBJECT), after=14)
    if letter.get("salutation"):
        par(letter["salutation"], after=8)
    for i, block in enumerate(letter["paragraphs"]):
        par(block, after=12, indent=letter["template"] == "schedule" and i == 0)

    if letter["template"] == "schedule":
        par(after=8)
        for barangay, when, vials in letter["rows"]:
            p = par(after=14)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(3.9), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.DASHES)
            p.add_run(f"{schedule_date(when)}\t {vials}vials Brgy.{barangay}")
        par(f"Total: {sum(r[2] for r in letter['rows'])} vials", bold=True, after=20)

    par(letter.get("closing") or "", after=4)
    sig = par()
    if signature_path:
        sig.add_run().add_picture(str(signature_path), height=Inches(0.6))
    par(letter["signatory_name"].upper(), bold=True)
    par(letter["signatory_position"])

    out = BytesIO()
    doc.save(out)
    return out.getvalue()
